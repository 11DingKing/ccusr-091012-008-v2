"""
仓库管理视图
"""
import logging
import io
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone
from django.db import transaction
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from apps.core.exceptions import BusinessException
from apps.personnel.qualification import (
    run_verified, QualificationInvalidError,
)
from .models import Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer,
    StockInSerializer, StockInCreateSerializer,
    StockOutSerializer, StockOutCreateSerializer,
    WarningSerializer,
    ApprovalSerializer, ApprovalDecisionSerializer,
)

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 货物、预警 ====================

class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = (
            Goods.objects.select_related(
                'variety', 'variety__category', 'variety__category__unit'
            ).filter(is_active=True).order_by('-created_at')
        )
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        goods = queryset[start:start + page_size]
        serializer = GoodsSerializer(goods, many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Warning.objects.select_related('goods').all().order_by('-created_at')
        is_read = request.query_params.get('is_read')
        if is_read in ('true', 'false'):
            queryset = queryset.filter(is_read=is_read == 'true')
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        warnings = queryset[start:start + page_size]
        serializer = WarningSerializer(warnings, many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


# ==================== 收件（入库）：核验“收件”资质 ====================

def _first_error(serializer):
    errors = serializer.errors
    first_error = list(errors.values())[0]
    if isinstance(first_error, (list, tuple)):
        first_error = first_error[0]
    return str(first_error)


class StockInListView(APIView):
    """收件（入库）记录列表 / 提交收件"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = (
            StockIn.objects.select_related('goods', 'operator', 'qualification_check')
            .all().order_by('-stock_in_time')
        )
        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        records = queryset[start:start + page_size]
        serializer = StockInSerializer(records, many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """收件：先核验值班员“收件”资质，通过后登记入库。"""
        serializer = StockInCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))
        data = serializer.validated_data

        def create_record(check):
            goods = Goods.objects.select_for_update().get(pk=data['goods'])
            record = StockIn.objects.create(
                goods=goods,
                operator=request.user,
                quantity=data['quantity'],
                batch_no=data.get('batch_no', ''),
                supplier=data.get('supplier', ''),
                remark=data.get('remark', ''),
                qualification_check=check,
            )
            goods.quantity = (goods.quantity or Decimal('0')) + data['quantity']
            goods.save(update_fields=['quantity', 'updated_at'])
            return record, goods

        # 资质核验未通过时抛出 QualificationInvalidError，由全局异常处理器返回403，
        # 阻断记录已独立固化，不会随业务事务回滚。
        record, goods = run_verified('receiving', request.user, create_record)

        logger.info(
            "Stock-in by %s: goods=%s quantity=%s check=%s",
            request.user.username, record.goods_id, record.quantity,
            record.qualification_check_id,
        )
        return success_response(data=StockInSerializer(record).data, message='收件成功')


# ==================== 出库申请、审批、放行 ====================

class StockOutListView(APIView):
    """出库申请列表 / 提交出库申请"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = (
            StockOut.objects.select_related('goods', 'operator', 'qualification_check')
            .all().order_by('-created_at')
        )
        status_filter = request.query_params.get('status')
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        records = queryset[start:start + page_size]
        serializer = StockOutSerializer(records, many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """提交出库申请：申请环节不核验资质，资质在审批、放行节点核验。"""
        serializer = StockOutCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer))
        data = serializer.validated_data

        goods = Goods.objects.get(pk=data['goods'])
        stock_out = StockOut.objects.create(
            goods=goods,
            operator=request.user,
            receiver=data['receiver'],
            receiver_dept=data.get('receiver_dept', ''),
            quantity=data['quantity'],
            remark=data.get('remark', ''),
            status='pending',
        )
        logger.info(
            "Stock-out application submitted by %s: goods=%s quantity=%s",
            request.user.username, goods.id, stock_out.quantity,
        )
        return success_response(data=StockOutSerializer(stock_out).data, message='申请已提交')


class ApprovalListView(APIView):
    """审批记录列表"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = (
            Approval.objects.select_related('stock_out', 'approver', 'qualification_check')
            .all().order_by('-created_at')
        )
        status_filter = request.query_params.get('status')
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size

        total = queryset.count()
        records = queryset[start:start + page_size]
        serializer = ApprovalSerializer(records, many=True)
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class ApprovalDecisionView(APIView):
    """审批决定：通过/拒绝均先核验审批人“审批”资质（拒绝也留痕）。"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk, decision):
        if decision not in ('approve', 'reject'):
            return error_response(message='无效的审批动作')
        serializer = ApprovalDecisionSerializer(data=request.data)
        serializer.is_valid(raise_exception=False)
        remark = serializer.validated_data.get('remark', '') if serializer.is_valid() else ''

        try:
            stock_out = StockOut.objects.select_related('goods').get(pk=pk)
        except StockOut.DoesNotExist:
            return error_response(message='出库申请不存在', code=404)

        if stock_out.status != 'pending':
            return error_response(message='该申请已审批，不能重复审批')

        new_status = 'approved' if decision == 'approve' else 'rejected'

        def create_approval(check):
            approval = Approval.objects.create(
                stock_out=stock_out,
                approver=request.user,
                status=new_status,
                remark=remark,
                qualification_check=check,
            )
            stock_out.status = new_status
            stock_out.save(update_fields=['status'])
            return approval

        approval = run_verified('approval', request.user, create_approval)

        logger.info(
            "Approval %s by %s: stock_out=%s check=%s",
            new_status, request.user.username, stock_out.id,
            approval.qualification_check_id,
        )
        return success_response(
            data=ApprovalSerializer(approval).data,
            message='审批通过' if decision == 'approve' else '已拒绝',
        )


class StockOutReleaseView(APIView):
    """放行：已审批通过的申请，放行时再次核验值班员“放行”资质并核减库存。"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            stock_out = StockOut.objects.select_related('goods').get(pk=pk)
        except StockOut.DoesNotExist:
            return error_response(message='出库申请不存在', code=404)

        if stock_out.status == 'completed':
            return error_response(message='该申请已放行完成')
        if stock_out.status != 'approved':
            return error_response(message='仅已审批通过的申请可以放行')

        try:
            def release(check):
                goods = Goods.objects.select_for_update().get(pk=stock_out.goods_id)
                if goods.quantity < stock_out.quantity:
                    raise BusinessException('库存不足，无法放行')
                goods.quantity = goods.quantity - stock_out.quantity
                goods.save(update_fields=['quantity', 'updated_at'])

                stock_out.status = 'completed'
                stock_out.stock_out_time = timezone.now()
                stock_out.qualification_check = check
                stock_out.last_block_reason = ''
                stock_out.last_block_at = None
                stock_out.save(update_fields=[
                    'status', 'stock_out_time', 'qualification_check',
                    'last_block_reason', 'last_block_at',
                ])
                return stock_out, goods

            stock_out, goods = run_verified('release', request.user, release)
        except QualificationInvalidError as exc:
            # 资质失效阻断：固化阻断原因，提示重新分配值班员
            stock_out.last_block_reason = exc.message
            stock_out.last_block_at = timezone.now()
            stock_out.save(update_fields=['last_block_reason', 'last_block_at'])
            logger.warning(
                "Release blocked for stock_out=%s by %s: %s",
                pk, request.user.username, exc.message,
            )
            return error_response(message=exc.message, code=403)

        logger.info(
            "Released by %s: stock_out=%s check=%s remaining=%s",
            request.user.username, stock_out.id,
            stock_out.qualification_check_id, goods.quantity,
        )
        return success_response(
            data=StockOutSerializer(stock_out).data, message='放行成功'
        )
